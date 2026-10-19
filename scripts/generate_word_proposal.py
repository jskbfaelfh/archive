import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_rtl(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    pPr.append(parse_xml(r'<w:bidi xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def set_run_font(run, font_name="Arial", size_pt=12, color_rgb=(0,0,0), bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'))
    rPr.append(parse_xml(r'<w:rtl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

def create_proposal_doc():
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Primary Color Palette
    PRIMARY = (15, 76, 129)     # Deep Academic Blue
    SECONDARY = (30, 41, 59)   # Slate Charcoal
    MUTED = (100, 116, 139)     # Subtle Gray
    ACCENT = (13, 148, 136)     # Teal Green

    # Header / University Banner
    p_header = doc.add_paragraph()
    set_rtl(p_header)
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("جامعة الزهراوي - رئاسة الجامعة\nشعبة شؤون التعليم والتحول الرقمي")
    set_run_font(r_uni, "Arial", 14, MUTED, bold=True)

    # Document Title
    p_title = doc.add_paragraph()
    set_rtl(p_title)
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("مذكرة عرض مشروع: منظومة الأرشفة الرقمية والحوكمة الأكاديمية")
    set_run_font(r_title, "Arial", 20, PRIMARY, bold=True)

    # Subtitle / Motto
    p_sub = doc.add_paragraph()
    set_rtl(p_sub)
    p_sub.paragraph_format.space_after = Pt(18)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("«التحول المؤسسي من فوضى قنوات التواصل غير الرسمية إلى منصة أكاديمية موحدة وموثوقة»")
    set_run_font(r_sub, "Arial", 12, ACCENT, bold=True)

    # Separator Line
    p_line = doc.add_paragraph()
    set_rtl(p_line)
    p_line.paragraph_format.space_after = Pt(14)
    r_line = p_line.add_run("_______________________________________________________________________________")
    set_run_font(r_line, "Arial", 9, (203, 213, 225))

    # SECTION 1: Executive Summary
    h1 = doc.add_paragraph()
    set_rtl(h1)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. الخلاصة التنفيذية (Executive Summary)")
    set_run_font(r_h1, "Arial", 15, PRIMARY, bold=True)

    p1 = doc.add_paragraph()
    set_rtl(p1)
    p1.paragraph_format.space_after = Pt(8)
    p1.paragraph_format.line_spacing = 1.25
    r_p1 = p1.add_run(
        "تهدف هذه المنظومة الرقمية إلى حل إحدى أكبر المعضلات الإدارية والتعليمية التي تواجه الكليات والأقسام "
        "حالياً، والمتمثلة في تشتت المحتوى الأكاديمي وضياعه عبر تطبيقات المراسلة الفورية غير المخصصة (تليغرام وواتساب). "
        "تم تصميم المنصة لتكون المستودع الرقمي الرسمي لجامعة الزهراوي، متوافقة حرفياً مع الهيكل الإداري العراقي "
        "(جامعة ⬅ كليات ⬅ أقسام ⬅ مراحل ⬅ شُعب)، بما يوفر للقيادة الأكاديمية أدوات قياس ورقابة لحظية كاملة."
    )
    set_run_font(r_p1, "Arial", 11.5, SECONDARY)

    # SECTION 2: Current Challenges vs Solution
    h2 = doc.add_paragraph()
    set_rtl(h2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("2. الواقع الحالي مقابل الحل الذي تقدمه المنظومة")
    set_run_font(r_h2, "Arial", 15, PRIMARY, bold=True)

    # Comparison Table
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Set table RTL
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        tblPr[0].append(parse_xml(r'<w:bidiVisual xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    headers = ["المعيار الأكاديمي", "الوضع التقليدي (التليغرام / الواتساب)", "الحل عبر المنظومة الرقمية للجامعة"]
    for i, title in enumerate(headers):
        cell = table.cell(0, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(cell.paragraphs[0])
        run = cell.paragraphs[0].add_run(title)
        set_run_font(run, "Arial", 11, (255, 255, 255), bold=True)
        # Background color
        shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="0F4C81"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    data = [
        ("حفظ الإرث والمقررات", "تضيع الملفات بمرور الأشهر، تُحذف الروابط، ولا يجد الطالب الجديد مراجع للسنوات السابقة.", "مستودع دائم موثق بصيغ PDF وعروض وفيديوهات يتبع اسم المقرر والقسم وسنة النشر بصورة غير قابلة للضياع."),
        ("حوكمة ورقابة العمادة والقسم", "انعدام الرؤية؛ لا يعلم رئيس القسم أو العميد متى تم رفع المحاضرة وهل التزم التدريسي بمفردات المقرر.", "لوحة تحكم لحظية تمنح العميد ورئيس القسم نسب إنجاز المقررات ومتابعة مواعيد نشر المحاضرات فورياً."),
        ("أمن البيانات والحسابات", "روابط عامة متاحة لأي شخص، وسهولة تزوير الصفحات أو انتحال هويات الطلبة.", "عزل أمني تام وصلاحيات هرمية (5 مستويات)؛ الطالب يرى مرحلته وشعبته فقط، والتدريسي مواده المسندة فقط.")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            set_rtl(cell.paragraphs[0])
            run = cell.paragraphs[0].add_run(text)
            set_run_font(run, "Arial", 10, SECONDARY, bold=(col_idx==0))
            if row_idx % 2 == 1:
                shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F8FAFC"/>')
                cell._tc.get_or_add_tcPr().append(shd)

    # SECTION 3: Architectural Pillars
    h3 = doc.add_paragraph()
    set_rtl(h3)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("3. الركائز الجوهرية للمنظومة (Core Capabilities)")
    set_run_font(r_h3, "Arial", 15, PRIMARY, bold=True)

    pillars = [
        ("أولاً: حوكمة الصلاحيات الخماسية (5-Tier Role Isolation):", 
         "توزيع الصلاحيات بدقة متناهية تشمل: المشرف العام (الرئاسة)، عمداء الكليات، رؤساء الأقسام، أساتذة المقررات، والطلبة. تمنع الصلاحيات تداخل التدريسيين في مقررات زملائهم، وتحصر الطالب ضمن شعبته ومواده حصراً."),
        ("ثانياً: سجل التدقيق الرقابي غير القابل للتلاعب (Immutable Audit Logs):",
         "توثيق كافة الأنشطة داخل المنظومة (رفع، تعديل، حذف، تحميل، استعادة) مسجلة بالدقيقة وهوية المستخدم ومسماه الأكاديمي، بما يضمن الشفافية والمساءلة الإدارية الكاملة."),
        ("ثالثاً: لوحة مؤشرات وتحليلات لصنّاع القرار (Decision Support KPIs):",
         "توليد إحصائيات بيانية فورية للعمادة ورئاسة الجامعة حول المقررات الأكثر تفاعلاً، عدد المشاهدات والتحميلات، ونسب التزام الأقسام بالخطة الدراسية."),
        ("رابعاً: استيراد البيانات الضخمة (Zero Setup Time):",
         "إمكانية تصدير واستيراد مئات الطلبة والمقررات دفعة واحدة عبر ملفات Excel وCSV، مع توليد فوري لكلمات مرور مؤقتة مشفرة وإجبار الطالب على تعيين كلمة مروره الخاصة عند أول دخول.")
    ]

    for title, desc in pillars:
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        r_t = p.add_run(f"• {title} ")
        set_run_font(r_t, "Arial", 11.5, PRIMARY, bold=True)
        r_d = p.add_run(desc)
        set_run_font(r_d, "Arial", 11, SECONDARY)

    # SECTION 4: Institutional ROI
    h4 = doc.add_paragraph()
    set_rtl(h4)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("4. العائد الأكاديمي والمؤسسي على الجامعة (Institutional ROI)")
    set_run_font(r_h4, "Arial", 15, PRIMARY, bold=True)

    rois = [
        "مطابقة متطلبات الاعتماد الأكاديمي الوطني والدولي ومعايير الجودة للتحول الرقمي.",
        "حماية الملكية الفكرية والبحوث والمحاضرات في مستودع رسمي مملوك للجامعة بنسبة 100%.",
        "توفير الوقت الإداري المستنزف في البحث عن المقررات وإعادة تداول المواد كل فصل دراسي.",
        "واجهة مستخدم عربية موحدة، فائقة البساطة والسرعة، تعمل على كافة الهواتف والحواسيب دون الحاجة لتدريب معقد."
    ]

    for item in rois:
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(3)
        r_chk = p.add_run("✔  ")
        set_run_font(r_chk, "Arial", 11, ACCENT, bold=True)
        r_txt = p.add_run(item)
        set_run_font(r_txt, "Arial", 11, SECONDARY)

    # Sign-off Box
    p_box = doc.add_paragraph()
    set_rtl(p_box)
    p_box.paragraph_format.space_before = Pt(20)
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_box = p_box.add_run("— تم إعداد هذه الوثيقة لغرض العرض والمصادقة على التطبيق الميداني —")
    set_run_font(r_box, "Arial", 11, MUTED, bold=True)

    output_path = os.path.abspath("Alzahrawi_System_Executive_Proposal.docx")
    doc.save(output_path)
    print("Document successfully created at:", output_path)

if __name__ == "__main__":
    create_proposal_doc()
